"""
File -> text parsing for ingestion. Runs inside the background worker, never in
the request handler, so a large PDF can't stall the API for other users.
"""

import html
import io
import logging
import os
import re
import tempfile
from typing import Optional

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {"txt", "md", "markdown", "csv", "tsv", "json", "jsonl", "xml", "yaml", "yml", "log", "rst", "ini", "sql"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {"pdf", "docx", "html", "htm"}

_docling = None


def extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _get_docling():
    global _docling
    if _docling is None:
        from app.services.docling_parser import DoclingParser
        _docling = DoclingParser()
    return _docling


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            text = raw.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _html_to_text(raw: bytes) -> str:
    text = _decode(raw)
    text = re.sub(r"(?is)<(script|style|noscript)[^>]*>.*?</\1>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h[1-6]|tr|section|article)>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = html.unescape(text)
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def _parse_pdf(raw: bytes) -> dict:
    docling = _get_docling()
    if docling.is_available():
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(raw)
            tmp_path = tmp.name
        try:
            parsed = docling.parse(tmp_path)
            content = "\n".join(b["content"] for b in parsed.get("text_blocks", []))
            content += "\n" + "\n".join(t["content"] for t in parsed.get("tables", []))
            return {
                "content": content.strip(),
                "page_count": parsed.get("page_count"),
                "parsing_method": parsed.get("parsing_method", "docling"),
                "parsed_chunks_json": {"chunks": docling.to_chunks(parsed)},
                "table_count": len(parsed.get("tables", [])),
                "heading_count": len(parsed.get("headings", [])),
            }
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return {
        "content": "\n".join(pages).strip(),
        "page_count": len(reader.pages),
        "parsing_method": "pypdf",
        "parsed_chunks_json": None,
    }


def _parse_docx(raw: bytes) -> dict:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise ValueError("DOCX support requires the python-docx package") from exc
    document = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return {"content": "\n".join(parts), "page_count": None, "parsing_method": "python-docx", "parsed_chunks_json": None}


def parse_file(filename: str, raw: bytes) -> dict:
    ext = extension(filename)
    if ext == "pdf":
        result = _parse_pdf(raw)
    elif ext == "docx":
        result = _parse_docx(raw)
    elif ext in ("html", "htm"):
        result = {"content": _html_to_text(raw), "page_count": None, "parsing_method": "html", "parsed_chunks_json": None}
    elif ext in TEXT_EXTENSIONS or not ext:
        result = {"content": _decode(raw), "page_count": None, "parsing_method": "basic", "parsed_chunks_json": None}
    else:
        raise ValueError(f"Unsupported file type: .{ext}")
    if not (result.get("content") or "").strip():
        raise ValueError("No extractable text found (scanned PDFs need OCR; install Docling for OCR support).")
    return result


def check_upload(filename: str, size: int, max_bytes: int) -> Optional[str]:
    """Return an error message for an unacceptable upload, else None."""
    ext = extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        return f"Unsupported file type '.{ext or '?'}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    if size == 0:
        return "File is empty"
    if size > max_bytes:
        return f"File exceeds the {max_bytes // (1024 * 1024)} MB limit"
    return None
